"""월별 손익 요약을 카톡/문자용 문구로 만들어주는 스크립트.

엑셀 수식 캐시값이 아니라 '일일기록' 원본 데이터를 직접 계산하므로,
Excel/LibreOffice로 파일을 열어 재계산하지 않아도 항상 정확한 값을 냅니다.

사용법: python3 src/automation/monthly_summary.py [연도] [월]
예:     python3 src/automation/monthly_summary.py 2026 1
인자를 생략하면 이번 달 기준으로 계산합니다.
"""

import datetime
import os
import sys

import openpyxl

FILE_PATH = os.path.join(
    os.path.dirname(__file__), "..", "..", "data", "vehicle", "5톤계란운송_손익관리.xlsx"
)


def load_params(wb):
    ws = wb["고정비설정"]
    values = {row[0].value: row[1].value for row in ws.iter_rows(min_row=3, max_row=7)}
    return {
        "unit_price": values.get("기준단가(원/판)") or 0,
        "driver_salary": values.get("월 기사월급(원)") or 0,
        "hipass_fixed": values.get("월 하이패스(원)") or 0,
        "insurance_annual": values.get("연 차량보험료(원)") or 0,
        "insurance_month": values.get("차량보험 납부월(1~12)") or 0,
    }


def to_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str):
        try:
            return datetime.datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def load_daily_records(wb):
    ws = wb["일일기록"]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        date = to_date(row[0])
        if date is None:
            continue
        records.append(
            {
                "date": date,
                "farm": row[1],
                "tray_count": row[2] or 0,
                "maintenance_cost": row[6] or 0,
                "hipass": row[7] or 0,
            }
        )
    return records


def summarize(records, params, year, month):
    month_records = [r for r in records if r["date"].year == year and r["date"].month == month]

    total_trays = sum(r["tray_count"] for r in month_records)
    revenue = total_trays * params["unit_price"]
    vat = round(revenue * 0.1)
    maintenance = sum(r["maintenance_cost"] for r in month_records)
    hipass = params["hipass_fixed"]
    driver_salary = params["driver_salary"]
    insurance = params["insurance_annual"] if month == params["insurance_month"] else 0
    total_expense = vat + maintenance + hipass + driver_salary + insurance
    net_profit = revenue - total_expense
    margin = (net_profit / revenue) if revenue else 0

    return {
        "total_trays": total_trays,
        "revenue": revenue,
        "vat": vat,
        "maintenance": maintenance,
        "hipass": hipass,
        "driver_salary": driver_salary,
        "insurance": insurance,
        "total_expense": total_expense,
        "net_profit": net_profit,
        "margin": margin,
    }


def won(n):
    return f"{round(n):,}원"


def build_message(year, month, s):
    lines = [
        f"[{year}년 {month}월 운송 정산]",
        f"총 판수: {s['total_trays']:,}판",
        f"매출액: {won(s['revenue'])}",
        f"부가세(10%): {won(s['vat'])}",
        f"차량유지비: {won(s['maintenance'])}",
        f"하이패스: {won(s['hipass'])}",
        f"기사월급: {won(s['driver_salary'])}",
        f"차량보험: {won(s['insurance'])}",
        f"총지출: {won(s['total_expense'])}",
        f"순이익: {won(s['net_profit'])}",
        f"이익률: {s['margin'] * 100:.1f}%",
        "",
        "금액 확인 부탁드립니다.",
    ]
    return "\n".join(lines)


def main():
    now = datetime.date.today()
    year = int(sys.argv[1]) if len(sys.argv) > 1 else now.year
    month = int(sys.argv[2]) if len(sys.argv) > 2 else now.month

    wb = openpyxl.load_workbook(FILE_PATH, data_only=False)
    params = load_params(wb)
    records = load_daily_records(wb)
    summary = summarize(records, params, year, month)
    print(build_message(year, month, summary))


if __name__ == "__main__":
    main()
